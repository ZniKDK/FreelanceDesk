"""Тесты экспорта в Excel и CSV."""

import csv
from datetime import date, datetime

from openpyxl import load_workbook

from freelancedesk.core.manager import OrderManager
from freelancedesk.core.storage import InMemoryStorage
from freelancedesk.demo import CLIENTS, seed
from freelancedesk.export import export_excel, export_payments_csv

TODAY = date(2026, 10, 15)


def demo_manager() -> OrderManager:
    manager = OrderManager(InMemoryStorage())
    seed(manager, TODAY)
    return manager


def test_excel_has_three_sheets_with_data(tmp_path):
    manager = demo_manager()
    path = tmp_path / "data.xlsx"
    export_excel(manager, path)

    book = load_workbook(path)
    assert book.sheetnames == ["Заказы", "Платежи", "Клиенты"]
    orders = book["Заказы"]
    assert orders["A1"].value == "Заказ" and orders["A1"].font.bold
    assert orders.max_row == len(manager.list_orders()) + 1
    # Цена — числом с форматом рублей, чтобы Excel мог считать
    assert isinstance(orders["C2"].value, (int, float))  # 15000.0 → 15000
    assert "₽" in orders["C2"].number_format
    payments = book["Платежи"]
    assert isinstance(payments["A2"].value, datetime)  # дата — датой
    assert book["Клиенты"].max_row == len(CLIENTS) + 1


def test_csv_for_period(tmp_path):
    manager = demo_manager()
    path = tmp_path / "payments.csv"
    export_payments_csv(manager, path, date(2026, 10, 1), TODAY)

    with open(path, encoding="utf-8-sig") as file:  # BOM для Excel
        rows = list(csv.reader(file, delimiter=";"))
    assert rows[0] == ["Дата", "Заказ", "Клиент", "Сумма", "Налог", "Чек"]
    expected = manager.payments_in_period(date(2026, 10, 1), TODAY)
    assert len(rows) - 1 == len(expected)
    # Деньги — с запятой, даты — ДД.ММ.ГГГГ
    assert "," in rows[1][3]
    assert rows[1][0].count(".") == 2
