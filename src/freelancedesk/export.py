"""Экспорт данных в Excel (.xlsx) и CSV.

Excel — все данные в одной книге на трёх листах: заказы, платежи,
клиенты. CSV — одна таблица (платежи за период) для сверки с «Мой
налог» или для бухгалтера. CSV пишется с разделителем «;» и меткой
UTF-8 (BOM): так русский Excel открывает файл без «кракозябр».
"""

import csv
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from freelancedesk.core.manager import OrderManager
from freelancedesk.core.models import ClientType, OrderStatus

STATUS_TEXT = {OrderStatus.NEW: "Новый", OrderStatus.IN_PROGRESS: "В работе",
               OrderStatus.DELIVERED: "Сдан", OrderStatus.CANCELLED: "Отменён"}
TYPE_TEXT = {ClientType.PERSON: "Физлицо", ClientType.COMPANY: "Юрлицо / ИП"}
MONEY_FORMAT = '#,##0.00 "₽"'
DATE_FORMAT = "DD.MM.YYYY"

PAYMENT_HEADERS = ["Дата", "Заказ", "Клиент", "Сумма", "Налог", "Чек"]


def order_rows(manager: OrderManager) -> tuple[list[str], list[list]]:
    names = {c.id: c.name for c in manager.list_clients()}
    money = manager.money_all()
    headers = ["Заказ", "Клиент", "Цена", "Получено", "Осталось",
               "Налог", "Статус", "Дедлайн", "Сдан", "Ссылка", "Описание"]
    rows = []
    for order in sorted(manager.list_orders(), key=lambda o: o.id):
        m = money[order.id]
        rows.append([order.title, names.get(order.client_id, ""),
                     order.amount, m.received, m.remaining, m.tax,
                     STATUS_TEXT[order.status], order.deadline,
                     order.delivered_on, order.link, order.description])
    return headers, rows


def payment_rows(manager: OrderManager, start: date | None = None,
                 end: date | None = None) -> list[list]:
    """Платежи (за период, если он задан) по дате."""
    orders = {o.id: o for o in manager.list_orders()}
    names = {c.id: c.name for c in manager.list_clients()}
    payments = (manager.payments_in_period(start, end)
                if start and end else
                manager.payments_in_period(date.min, date.max))
    rows = []
    for payment in payments:
        order = orders.get(payment.order_id)
        rows.append([payment.paid_on, order.title if order else "",
                     names.get(order.client_id, "") if order else "",
                     payment.amount, manager.payment_tax(payment),
                     "выбит" if payment.receipt_issued else "не выбит"])
    return rows


def client_rows(manager: OrderManager) -> tuple[list[str], list[list]]:
    stats = manager.client_stats()
    headers = ["Имя", "Тип", "Площадка", "Почта", "Телефон", "Мессенджер",
               "Заказов", "Получено", "Заметка"]
    rows = []
    for client in manager.list_clients():
        s = stats.get(client.id)
        messenger = " ".join(filter(None, (client.messenger,
                                           client.messenger_app
                                           if client.messenger else "")))
        rows.append([client.name, TYPE_TEXT[client.client_type],
                     client.platform, client.email, client.phone, messenger,
                     s.orders if s else 0, s.income if s else Decimal("0"),
                     client.note])
    return headers, rows


def _write_sheet(sheet, headers: list[str], rows: list[list]) -> None:
    """Лист с жирной шапкой, форматом денег и дат, шириной по тексту."""
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E7EFFD")
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        # Decimal openpyxl пишет числом — Excel сможет посчитать сумму
        sheet.append([float(v) if isinstance(v, Decimal) else v
                      for v in row])
    for column, header in enumerate(headers, start=1):
        letter = get_column_letter(column)
        values = [header] + [row[column - 1] for row in rows]
        for cell in sheet[letter][1:]:
            if isinstance(rows[0][column - 1] if rows else None, Decimal):
                cell.number_format = MONEY_FORMAT
            elif isinstance(cell.value, date):
                cell.number_format = DATE_FORMAT
        width = max(len(str(v)) if v is not None else 0 for v in values)
        sheet.column_dimensions[letter].width = min(max(width + 2, 10), 50)
    sheet.freeze_panes = "A2"  # шапка не уезжает при прокрутке


def export_excel(manager: OrderManager, path: Path) -> None:
    """Все данные в одну книгу Excel: заказы, платежи, клиенты."""
    book = Workbook()
    orders_sheet = book.active
    orders_sheet.title = "Заказы"
    _write_sheet(orders_sheet, *order_rows(manager))
    _write_sheet(book.create_sheet("Платежи"), PAYMENT_HEADERS,
                 payment_rows(manager))
    _write_sheet(book.create_sheet("Клиенты"), *client_rows(manager))
    book.save(path)


def export_payments_csv(manager: OrderManager, path: Path, start: date,
                        end: date) -> None:
    """Платежи за период в CSV (для «Мой налог» или бухгалтера)."""
    with open(path, "w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file, delimiter=";")
        writer.writerow(PAYMENT_HEADERS)
        for row in payment_rows(manager, start, end):
            writer.writerow([
                v.strftime("%d.%m.%Y") if isinstance(v, date)
                # В русском Excel дробная часть — через запятую
                else f"{v:.2f}".replace(".", ",") if isinstance(v, Decimal)
                else v
                for v in row])
